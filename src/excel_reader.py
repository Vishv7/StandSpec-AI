"""
Excel reader and standard-number normalizer for BIS department Excel files.
Phase 1 of the BIS Standards Data Collector.

Uses openpyxl directly (no pandas) to avoid DLL issues on restricted systems.

BIS Excel format:
  Row 1: Metadata row (department name, generated date) — merged cells
  Row 2: Actual column headers (Sl#, Standard Number, Date of Publish, Title, ...)
  Row 3+: Data rows
  Trailing: many empty rows from the original full export
"""

import re
import csv
import json
import hashlib
import logging
from pathlib import Path
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

# ── Column name mapping ──
# BIS exports use slightly different headers across departments.
COLUMN_ALIASES = {
    "standard_number": [
        "standard number", "standardnumber", "std number", "std no",
        "is number", "is no", "standard no", "standard_number",
        "is number (latest)", "standard number (latest)",
    ],
    "title": [
        "title", "standard title", "standardtitle", "name",
        "title of standard", "title of is",
    ],
    "date_of_publish": [
        "date of publish", "dateofpublish", "published date",
        "publish date", "date of publication", "published on",
        "date_of_publish",
    ],
    "type_of_standard": [
        "type of standard", "typeofstandard", "standard type",
        "type of standard", "type_of_standard",
    ],
    "degree_of_equivalence": [
        "degree of equivalence", "degreeofequivalence",
        "equivalence", "degree_of_equivalence",
    ],
    "sl_number": [
        "sl#", "sl no", "sl.", "serial", "s.no", "sr no", "sno", "s no",
    ],
}


def _find_column_index(headers: list[str], aliases: list[str]) -> int | None:
    """Find a column index matching any of the given aliases (case-insensitive)."""
    lower_headers = {}
    for i, h in enumerate(headers):
        if h:
            lower_headers[h.strip().lower()] = i
    
    for alias in aliases:
        alias_lower = alias.strip().lower()
        if alias_lower in lower_headers:
            return lower_headers[alias_lower]
    # Partial/substring match as fallback
    for alias in aliases:
        alias_lower = alias.strip().lower()
        for header_lower, idx in lower_headers.items():
            if alias_lower in header_lower or header_lower in alias_lower:
                return idx
    return None


def map_columns(headers: list[str]) -> dict[str, int | None]:
    """Map expected field names to actual column indices."""
    mapping = {}
    for field, aliases in COLUMN_ALIASES.items():
        idx = _find_column_index(headers, aliases)
        mapping[field] = idx
        if idx is not None:
            logger.info(f"Mapped '{field}' -> column {idx} ('{headers[idx]}')")
        else:
            if field != "sl_number":  # sl_number is optional
                logger.warning(f"Could not find column for '{field}' in {headers}")
    return mapping


def normalize_standard_number(raw: str) -> dict:
    """
    Normalize a BIS standard number into a consistent internal representation.
    Preserves the original string alongside the normalized form.
    
    Returns a dict with:
        - original: the raw string as-is
        - is_number: normalized form without year (e.g. "IS 302 (Part 2/Sec 16)")
        - year: extracted year as int or None
        - recognized: bool, whether the format was recognized
    """
    if not raw or not isinstance(raw, str):
        return {
            "original": str(raw) if raw is not None else "",
            "is_number": None,
            "year": None,
            "recognized": False,
        }
    
    raw = raw.strip()
    
    # Pattern: IS <number> [(<Part X>[/<Sec Y>])] [:<year>]
    # Also handle SP, IS/IEC, IS/ISO patterns
    pattern = re.compile(
        r'^(IS|SP|IS/IEC|IS/ISO)\s+'             # prefix
        r'(\d+)'                                   # base number
        r'(?:\s*\(\s*Part\s*(\d+)'                 # optional Part number
        r'(?:\s*/\s*Sec\s*(\d+))?\s*\))?'          # optional Section number  
        r'(?:\s*:\s*(\d{4}))?'                     # optional :year
        r'\s*$',
        re.IGNORECASE
    )
    
    m = pattern.match(raw)
    if not m:
        return {
            "original": raw,
            "is_number": None,
            "year": None,
            "recognized": False,
        }
    
    prefix = m.group(1).upper()
    base_num = m.group(2)
    part = m.group(3)
    section = m.group(4)
    year_str = m.group(5)
    
    # Build normalized is_number (without year)
    is_number = f"{prefix} {base_num}"
    if part:
        if section:
            is_number += f" (Part {part}/Sec {section})"
        else:
            is_number += f" (Part {part})"
    
    year = int(year_str) if year_str else None
    
    return {
        "original": raw,
        "is_number": is_number,
        "year": year,
        "recognized": True,
    }


def _detect_header_row(ws) -> int:
    """
    Auto-detect which row contains the actual column headers.
    BIS files have a metadata row first (department name, date), 
    then the real headers (Standard Number, Title, etc.).
    
    Returns 1-indexed row number.
    """
    for row_num in range(1, min(10, ws.max_row + 1)):
        values = [ws.cell(row=row_num, column=col).value for col in range(1, ws.max_column + 1)]
        values_str = [str(v).strip().lower() if v else "" for v in values]
        
        # Look for "standard number" or "title" in this row — that's the header
        if any("standard number" in v or "standard no" in v for v in values_str):
            return row_num
        if any("title" in v and "type" in "".join(values_str) for v in values_str):
            return row_num
    
    # Default: assume row 1 is headers
    return 1


def read_excel(file_path: str | Path) -> tuple[list[str], list[list], dict[str, int | None], dict]:
    """
    Read a BIS department Excel file using openpyxl directly.
    Auto-detects the header row (BIS files have a metadata row first).
    
    Returns:
        - headers: list of column header strings
        - rows: list of row data (each row is a list of cell values), only non-empty rows
        - col_map: maps our expected field names to column indices
        - metadata: dict with department name, generated date from the metadata row
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")
    
    logger.info(f"Reading Excel file: {file_path}")
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    
    max_col = ws.max_column
    max_row = ws.max_row
    logger.info(f"Sheet: '{ws.title}', max_row={max_row}, max_col={max_col}")
    
    # Detect header row
    header_row_num = _detect_header_row(ws)
    logger.info(f"Header row detected at row {header_row_num}")
    
    # Extract metadata from row(s) before header
    metadata = {
        "sheet_name": ws.title,
        "header_row_num": header_row_num,
    }
    for row_num in range(1, header_row_num):
        for col in range(1, max_col + 1):
            val = ws.cell(row=row_num, column=col).value
            if val:
                val_str = str(val).strip()
                if "department" in val_str.lower() or "etd" in val_str.lower():
                    metadata["department_name"] = val_str
                if "generated" in val_str.lower():
                    metadata["generated_on"] = val_str
    
    # Read headers
    headers = []
    for col in range(1, max_col + 1):
        val = ws.cell(row=header_row_num, column=col).value
        headers.append(str(val).strip() if val else "")
    
    logger.info(f"Headers: {headers}")
    
    # Read data rows (skip empty rows)
    data_rows = []
    row_numbers = []
    for row_num in range(header_row_num + 1, max_row + 1):
        row_values = []
        for col in range(1, max_col + 1):
            row_values.append(ws.cell(row=row_num, column=col).value)
        
        # Skip entirely empty rows
        if all(v is None for v in row_values):
            continue
        
        # Skip rows where the standard number column is empty
        # (trailing metadata/empty rows in BIS files)
        std_idx = _find_column_index(headers, COLUMN_ALIASES["standard_number"])
        if std_idx is not None:
            std_val = row_values[std_idx] if std_idx < len(row_values) else None
            if std_val is None or str(std_val).strip() == "":
                continue
        
        data_rows.append(row_values)
        row_numbers.append(row_num)
    
    wb.close()
    metadata["row_numbers"] = row_numbers
    
    logger.info(f"Loaded {len(data_rows)} data rows (skipped {max_row - header_row_num - len(data_rows)} empty/trailing rows)")
    
    col_map = map_columns(headers)
    
    return headers, data_rows, col_map, metadata


def _safe_cell_value(row: list, idx: int | None) -> str | None:
    """Safely extract a cell value from a row by index."""
    if idx is None or idx >= len(row):
        return None
    val = row[idx]
    if val is None:
        return None
    return str(val).strip()


def process_excel(file_path: str | Path, reports_dir: str | Path = "data/reports") -> tuple[list[dict], list[str], list[list]]:
    """
    Read the Excel file, normalize all standard numbers, and flag unrecognized formats.
    
    Returns:
        - standards_list: list of dicts with original, is_number, year, title, raw_excel, etc.
        - headers: list of column header strings (for later Excel enrichment)
        - raw_rows: list of raw row data (for later Excel enrichment)
    """
    reports_dir = Path(reports_dir)
    reports_dir.mkdir(parents=True, exist_ok=True)
    
    headers, data_rows, col_map, metadata = read_excel(file_path)
    
    logger.info(f"File metadata: {metadata}")
    
    std_idx = col_map.get("standard_number")
    if std_idx is None:
        raise ValueError(f"Cannot find 'Standard Number' column in {headers}")
    
    title_idx = col_map.get("title")
    date_idx = col_map.get("date_of_publish")
    type_idx = col_map.get("type_of_standard")
    equiv_idx = col_map.get("degree_of_equivalence")
    
    standards = []
    unrecognized = []
    
    for row_idx, row in enumerate(data_rows):
        raw_std = _safe_cell_value(row, std_idx) or ""
        if not raw_std:
            continue
        
        normalized = normalize_standard_number(raw_std)
        physical_row = metadata.get("row_numbers", [])[row_idx] if "row_numbers" in metadata else (metadata.get("header_row_num", 1) + 1 + row_idx)
        sheet_name = metadata.get("sheet_name", "Sheet1")
        
        raw_columns = {h: (str(v) if v is not None else None) for h, v in zip(headers, row) if h}
        row_sha256 = hashlib.sha256(json.dumps(raw_columns, sort_keys=True).encode("utf-8")).hexdigest()
        
        raw_excel_payload = {
            "sheet_name": sheet_name,
            "row_number": physical_row,
            "row_sha256": row_sha256,
            "columns": raw_columns,
        }
        
        record = {
            "row_index": row_idx,
            "sheet_name": sheet_name,
            "physical_row_number": physical_row,
            "original_standard_number": normalized["original"],
            "is_number": normalized["is_number"],
            "year": normalized["year"],
            "recognized": normalized["recognized"],
            "title": _safe_cell_value(row, title_idx),
            "date_of_publish": _safe_cell_value(row, date_idx),
            "type_of_standard": _safe_cell_value(row, type_idx),
            "degree_of_equivalence": _safe_cell_value(row, equiv_idx),
            "raw_excel": raw_excel_payload,
        }
        
        standards.append(record)
        
        if not normalized["recognized"]:
            unrecognized.append({
                "row_index": row_idx,
                "original_standard_number": normalized["original"],
                "reason": "Format not recognized by normalizer",
            })
    
    # Write unrecognized formats report
    if unrecognized:
        unrec_path = reports_dir / "unrecognized_formats.csv"
        with open(unrec_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=["row_index", "original_standard_number", "reason"])
            writer.writeheader()
            writer.writerows(unrecognized)
        logger.warning(f"Found {len(unrecognized)} unrecognized standard number formats -> {unrec_path}")
    else:
        logger.info("All standard numbers recognized successfully.")
    
    logger.info(f"Processed {len(standards)} standards: {sum(1 for s in standards if s['recognized'])} recognized, {len(unrecognized)} unrecognized")
    
    return standards, headers, data_rows
