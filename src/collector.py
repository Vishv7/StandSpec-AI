"""
Main collector orchestrator -- runs the full pipeline end-to-end.
Phase 5 of the BIS Standards Data Collector — V1.2 Hardened.

Conforms to V1.2 Collector Data Contract:
- Preserves raw input Excel data including all original columns (A11)
- Computes canonical standard identity with year-aware verification (A8)
- Source URL provenance: verified_source_url vs candidate_urls (A9)
- Non-overlapping status accounting: primary_status + quality.manual_review_required (A10)
- MATCH_VERSION_UNCERTAIN pages flagged — Scope/References NOT treated as verified evidence
- Run manifest with parser/collector/schema versions and timestamps (A12)
- Atomic state file writes (A17)
- Idempotent output keyed by canonical standard designation
- Zero duplicates on repeated runs
- Emits standards.jsonl, references.jsonl, source_availability.jsonl/csv, and collection_report.json
- Supports offline cached reparsing mode (--parse-cache)
"""

import json
import time
import csv
import hashlib
import logging
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from src import __version__, PARSER_VERSION
from src.version import COLLECTOR_SCHEMA_VERSION as SCHEMA_VERSION
from src.excel_reader import process_excel, normalize_standard_number
from src.preview_id import (
    parse_standard_identity,
    standard_number_to_preview_id,
    standard_number_to_preview_id_candidates,
    get_preview_url,
)
from src.fetcher import PreviewFetcher
from src.parser import parse_preview_html, _normalize_title
from src.validators import validate_standard_v1_2, validate_reference_v1_2

logger = logging.getLogger(__name__)


def _get_git_commit_metadata() -> tuple[str | None, str]:
    """Get the current git commit hash and provenance source."""
    try:
        result = subprocess.run(
            ["git", "rev-parse", "--short", "HEAD"],
            capture_output=True, text=True, timeout=5
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip(), "GIT_HEAD"
    except Exception:
        pass
    return None, "NO_GIT_METADATA"


def _get_git_commit() -> str | None:
    """Get the current git commit hash, if available."""
    commit, _ = _get_git_commit_metadata()
    return commit


def _compute_file_sha256(file_path: Path) -> str:
    """Compute SHA-256 hash of a file."""
    h = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


class Collector:
    """
    End-to-end BIS Standards Data Collector adhering to the V1.2 Data Contract.
    
    Resumable & Idempotent: If interrupted or re-run, canonical records are updated
    by standard designation without duplicate record accumulation.
    """
    
    def __init__(
        self,
        input_excel: str | Path,
        department: str,
        output_dir: str | Path = None,
        raw_dir: str | Path = "data/raw/preview_html",
        reports_dir: str | Path = "data/reports",
        logs_dir: str | Path = "data/logs",
        parse_cache_only: bool = False,
        force_reparse: bool = False,
        incremental: bool = False,
    ):
        self.input_excel = Path(input_excel)
        self.department = department.upper()
        self.output_dir = Path(output_dir) if output_dir else Path(f"data/processed/{self.department}")
        self.raw_dir = Path(raw_dir)
        self.reports_dir = Path(reports_dir)
        self.logs_dir = Path(logs_dir)
        self.parse_cache_only = parse_cache_only
        self.force_reparse = force_reparse
        self.incremental = incremental
        
        # Ensure directories exist
        for d in [self.output_dir, self.raw_dir, self.reports_dir, self.logs_dir]:
            d.mkdir(parents=True, exist_ok=True)
            
        # State file for resumability
        self.state_file = self.output_dir / "state.json"
        self.state = self._load_state()
        
        # Output file paths
        self.standards_jsonl = self.output_dir / f"{self.department}_standards.jsonl"
        self.references_jsonl = self.output_dir / f"{self.department}_references.jsonl"
        self.availability_jsonl = self.output_dir / f"{self.department}_source_availability.jsonl"
        self.availability_csv = self.output_dir / f"{self.department}_source_availability.csv"
        self.enriched_excel = self.output_dir / f"{self.department}_enriched.xlsx"
        self.report_path = self.reports_dir / f"{self.department}_collection_report.json"
        
        # In-memory stores for guaranteed idempotency
        self.records_by_id = {}
        self.references_by_key = {}
        self.availability_by_id = {}
        
        # Preload existing records if present
        self._preload_existing_records()
        
        # Fetcher
        self.fetcher = PreviewFetcher(cache_dir=self.raw_dir)
        
        # Statistics
        self.stats = {
            "success": 0,
            "partial_success": 0,
            "failed": 0,
            "review_required": 0,
            "skipped_cached": 0,
            "reparsed": 0,
        }
        
    def _load_state(self) -> dict:
        """Load processing state from disk."""
        if self.state_file.exists():
            try:
                with open(self.state_file, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                logger.warning("Corrupted state file, starting fresh")
        return {
            "completed": {},
            "first_run_started_at": None,
            "last_run_started_at": None,
            "run_started_at": None,
        }
        
    def _save_state(self):
        """A17: Persist processing state to disk atomically."""
        tmp = self.state_file.with_suffix(".json.tmp")
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.state, f, ensure_ascii=False, indent=2, default=str)
        tmp.replace(self.state_file)
            
    def _preload_existing_records(self):
        """Preload existing JSONL records into memory to guarantee idempotency across reruns."""
        if self.standards_jsonl.exists():
            try:
                with open(self.standards_jsonl, "r", encoding="utf-8") as f:
                    for line in f:
                        line = line.strip()
                        if line:
                            rec = json.loads(line)
                            # Identify canonical designation
                            desig = rec.get("identity", {}).get("standard_designation")
                            if not desig:
                                raw = rec.get("input", {}).get("raw_standard_number") or rec.get("is_number")
                                year = rec.get("year")
                                if raw and year and f":{year}" not in str(raw):
                                    raw = f"{raw}:{year}"
                                ident = parse_standard_identity(str(raw)) if raw else None
                                desig = ident["standard_designation"] if ident else raw
                                
                            # Only store V1.1+ records or let new V1.2 records cleanly overwrite
                            if desig and rec.get("record_version") in ("1.1", "1.2"):
                                self.records_by_id[desig] = rec
            except Exception as e:
                logger.warning(f"Could not preload existing standards: {e}")
                
        if self.references_jsonl.exists():
            try:
                with open(self.references_jsonl, "r", encoding="utf-8") as f:
                    for line in f:
                        line = line.strip()
                        if line:
                            ref = json.loads(line)
                            # Only preload V1.1+ structured references
                            if isinstance(ref.get("target"), dict):
                                src = ref.get("source_standard")
                                tgt = ref["target"].get("designation")
                                rel = ref.get("relationship")
                                if src and tgt and rel:
                                    self.references_by_key[(src, tgt, rel)] = ref
            except Exception as e:
                logger.warning(f"Could not preload existing references: {e}")

    def run(self) -> dict:
        """
        Run the full collection pipeline end-to-end.
        """
        run_start = datetime.now(timezone.utc)
        if not self.state.get("run_started_at"):
            self.state["run_started_at"] = run_start.isoformat()
            self._save_state()
            
        logger.info("=" * 60)
        logger.info(f"BIS Standards Data Collector V1.2 -- {self.department}")
        logger.info(f"Input: {self.input_excel}")
        logger.info(f"Output: {self.output_dir}")
        logger.info(f"Mode: {'OFFLINE CACHED PARSE' if self.parse_cache_only else 'ONLINE FETCH & PARSE'}")
        logger.info("=" * 60)
        
        # Phase 1: Read Excel
        standards, headers, raw_rows = process_excel(self.input_excel, self.reports_dir)
        logger.info(f"Loaded {len(standards)} standards from Excel")
        
        # Track first_run_started_at (immutable) vs last_run_started_at (updated each run)
        if not self.state.get("first_run_started_at"):
            self.state["first_run_started_at"] = self.state.get("run_started_at") or run_start.isoformat()
        self.state["last_run_started_at"] = run_start.isoformat()
        self.state["run_started_at"] = self.state["first_run_started_at"]  # for backward compatibility
        
        # A12: Build run manifest and check for version changes
        input_sha256 = _compute_file_sha256(self.input_excel)
        git_commit, commit_source = _get_git_commit_metadata()
        run_manifest = {
            "input_sha256": input_sha256,
            "source_snapshot_sha256": input_sha256,
            "parser_version": PARSER_VERSION,
            "collector_version": __version__,
            "schema_version": SCHEMA_VERSION,
            "first_run_started_at": self.state["first_run_started_at"],
            "last_run_started_at": self.state["last_run_started_at"],
            "source_collection_timestamp": run_start.isoformat(),
            "code_commit": git_commit,
            "code_version_source": commit_source,
        }
        
        prev_manifest = self.state.get("run_manifest")
        if prev_manifest:
            changes = []
            if prev_manifest.get("parser_version") != PARSER_VERSION:
                changes.append(f"parser: {prev_manifest.get('parser_version')} -> {PARSER_VERSION}")
            if prev_manifest.get("schema_version") != SCHEMA_VERSION:
                changes.append(f"schema: {prev_manifest.get('schema_version')} -> {SCHEMA_VERSION}")
            if prev_manifest.get("input_sha256") != input_sha256:
                changes.append("input file changed")
            if changes:
                logger.warning(f"Version/input changes detected: {', '.join(changes)}")
                logger.warning("Consider using --force-reparse to re-parse all cached HTML")
        
        self.state["run_manifest"] = run_manifest
        self._save_state()
        
        # Process each standard
        total = len(standards)
        active_canonical_ids = set()
        
        # Pre-compute active canonical IDs for snapshot mode enforcement (Gate P0-3)
        for std in standards:
            rnum = std["original_standard_number"]
            idt = parse_standard_identity(rnum)
            cid = idt["standard_designation"] if idt else rnum
            active_canonical_ids.add(cid)
            
        for i, std in enumerate(standards, 1):
            raw_num = std["original_standard_number"]
            ident = parse_standard_identity(raw_num)
            canonical_id = ident["standard_designation"] if ident else raw_num
            
            prior_rec = self.records_by_id.get(canonical_id)
            prior_status = self.state.get("completed", {}).get(canonical_id, {}).get("status")
            
            # Gate P0-2: Invalidation checks
            should_reparse = self.force_reparse
            if not should_reparse and prior_rec:
                prior_prov = prior_rec.get("provenance", {})
                prior_raw = prior_rec.get("raw_excel", {})
                current_raw = std.get("raw_excel", {})
                
                if prior_prov.get("parser_version") != PARSER_VERSION:
                    should_reparse = True
                elif prior_prov.get("schema_version") != SCHEMA_VERSION:
                    should_reparse = True
                elif current_raw.get("row_sha256") and prior_raw.get("row_sha256") != current_raw.get("row_sha256"):
                    should_reparse = True
            
            # Check resumability: skip if already successfully processed and not invalidated
            if not should_reparse and prior_status == "success" and not self.parse_cache_only and canonical_id in self.records_by_id:
                self.stats["skipped_cached"] += 1
                logger.debug(f"[{i}/{total}] Skipping already completed (valid cache): {canonical_id}")
                continue
                
            if should_reparse and canonical_id in self.records_by_id:
                self.stats["reparsed"] += 1
                logger.info(f"[{i}/{total}] Re-parsing invalidated/forced record: {canonical_id}")
            else:
                logger.info(f"[{i}/{total}] Processing: {canonical_id}")
                
            record, refs, avail = self._process_one_standard(std, ident, prior_rec=prior_rec)
            
            # Gate P0-1: Strict Schema Validation in Pipeline
            std_errors = validate_standard_v1_2(record)
            if std_errors:
                err_msg = f"Schema validation error on {canonical_id}: {'; '.join(std_errors)}"
                logger.error(err_msg)
                record["quality"]["errors"].append(err_msg)
                record["quality"]["manual_review_required"] = True
                record["primary_status"] = "failed"
                
            for ref in refs:
                ref_errors = validate_reference_v1_2(ref)
                if ref_errors:
                    logger.error(f"Reference validation error on {canonical_id} -> {ref.get('target', {}).get('designation')}: {'; '.join(ref_errors)}")
            
            # Store in idempotent dictionaries
            self.records_by_id[canonical_id] = record
            for ref in refs:
                src = ref["source_standard"]
                tgt = ref["target"]["designation"]
                rel = ref["relationship"]
                self.references_by_key[(src, tgt, rel)] = ref
            self.availability_by_id[canonical_id] = avail
            
            # A10: Update stats using primary_status only (no overlap)
            status = record["primary_status"]
            if status == "success":
                self.stats["success"] += 1
            elif status == "partial_success":
                self.stats["partial_success"] += 1
            else:
                self.stats["failed"] += 1
            if record["quality"]["manual_review_required"]:
                self.stats["review_required"] += 1
                
            self.state.setdefault("completed", {})[canonical_id] = {
                "status": status,
                "completed_at": datetime.now(timezone.utc).isoformat(),
            }
            
            # Periodic state save
            if i % 25 == 0:
                self._save_state()
                self._consolidate_output_files(headers, raw_rows, active_ids=active_canonical_ids)
                
        # Final atomic consolidation of output files
        self._consolidate_output_files(headers, raw_rows, active_ids=active_canonical_ids)
        self._save_state()
        
        run_end = datetime.now(timezone.utc)
        report = self._generate_report(standards, run_start, run_end)
        with open(self.report_path, "w", encoding="utf-8") as f:
            json.dump(report, f, indent=2, ensure_ascii=False, default=str)
        logger.info(f"Run report written to {self.report_path}")
        
        self.fetcher.close()
        return report
        
    def _process_one_standard(self, std: dict, ident: dict | None, prior_rec: dict | None = None) -> tuple[dict, list[dict], dict]:
        """
        Process a single standard according to the V1.2 Data Contract.
        Returns: (standard_record, list_of_references, availability_record)
        """
        raw_num = std["original_standard_number"]
        canonical_id = ident["standard_designation"] if ident else raw_num
        candidates = standard_number_to_preview_id_candidates(
            raw_num, publish_date=std.get("date_of_publish")
        )
        
        warnings = []
        errors = []
        review_required = False
        
        # Base identity dict
        identity_dict = ident if ident else {
            "standard_designation": raw_num,
            "family": "UNKNOWN",
            "base_number": raw_num,
            "part": None,
            "section": None,
            "year": str(std.get("year")) if std.get("year") else None,
        }
        
        if not candidates:
            warnings.append("Could not generate candidate preview IDs")
            review_required = True
            fetch_result = {
                "requested_preview_id": None,
                "matched_preview_id": None,
                "preview_url": None,
                "attempted_urls": [],
                "http_status": None,
                "fetch_status": "failed",
                "identity_match": "unknown",
                "page_identity_evidence": None,
                "html": None,
                "cache_path": None,
                "content_valid": False,
                "error": "Unrecognized standard number format",
            }
        elif self.parse_cache_only:
            # Offline cache mode: check local cache directly without network
            fetch_result = self._fetch_from_cache_only(candidates, ident)
        else:
            fetch_result = self.fetcher.fetch_with_fallback(candidates, requested_identity=ident)
            
        html = fetch_result.get("html")
        identity_match = fetch_result.get("identity_match", "unknown")
        
        # A8+A9: Determine if content should be treated as verified evidence
        # Gate P0-5: Only match_version_verified pages are treated as verified evidence
        is_verified_evidence = (identity_match in ("match_version_verified", "match"))
        
        if html and fetch_result.get("content_valid"):
            try:
                parsed = parse_preview_html(html, source_standard=canonical_id)
            except Exception as e:
                logger.error(f"Parse error for {canonical_id}: {e}")
                errors.append(f"Parser error: {str(e)}")
                review_required = True
                parsed = {
                    "header_text": None,
                    "title": None,
                    "title_source": "excel",
                    "ics_raw": None,
                    "ics_codes": [],
                    "ics_code": None,
                    "committee": None,
                    "scope": {"text": None, "status": "not_verified", "reason": "parse_error", "extraction_method": None, "scope_source_section": None},
                    "national_foreword": {"text": None, "status": "not_verified", "reason": "parse_error", "extraction_method": None},
                    "notes": {"items": [], "status": "not_verified", "reason": "parse_error", "extraction_method": None},
                    "formal_references": [],
                    "prose_references": [],
                    "parse_warnings": [f"Parser error: {str(e)}"],
                }
        else:
            parsed = {
                "header_text": None,
                "title": None,
                "title_source": "excel",
                "ics_raw": None,
                "ics_codes": [],
                "ics_code": None,
                "committee": None,
                "scope": {"text": None, "status": "not_verified", "reason": "preview_unavailable", "extraction_method": None, "scope_source_section": None},
                "national_foreword": {"text": None, "status": "not_verified", "reason": "preview_unavailable", "extraction_method": None},
                "notes": {"items": [], "status": "not_verified", "reason": "preview_unavailable", "extraction_method": None},
                "formal_references": [],
                "prose_references": [],
                "parse_warnings": [],
            }
            if fetch_result.get("fetch_status") == "failed":
                errors.append(fetch_result.get("error") or "Preview fetch failed")
                
        # Merge references and deduplicate
        all_refs = []
        seen_ref_keys = set()
        for ref in (parsed["formal_references"] + parsed["prose_references"]):
            ref_key = (ref["target"]["designation"], ref["relationship"])
            if ref_key not in seen_ref_keys:
                seen_ref_keys.add(ref_key)
                all_refs.append(ref)
                
        # Check review required conditions
        if identity_match == "unknown":
            review_required = True
            warnings.append("Preview page identity is UNKNOWN")
        if identity_match in ("match_version_uncertain", "match_identity_only_version_uncertain"):
            review_required = True
            evidence = fetch_result.get("page_identity_evidence", {}) or {}
            warnings.append(
                f"Version uncertain ({identity_match}): requested year '{evidence.get('requested_year')}', "
                f"page year '{evidence.get('page_year')}'. "
                f"Scope/References NOT treated as verified evidence."
            )
        if fetch_result.get("fetch_status") == "review_required":
            review_required = True
        if parsed.get("parse_warnings"):
            warnings.extend(parsed["parse_warnings"])
            review_required = True
            
        title_val = parsed.get("title") or _normalize_title(std.get("title"))
        
        # A10: Non-overlapping status accounting
        # primary_status is ONLY success/partial_success/failed
        # quality.manual_review_required is a separate boolean
        if fetch_result.get("fetch_status") == "success" and parsed["scope"]["status"] in ("present", "fallback"):
            primary_status = "success"
        elif fetch_result.get("content_valid"):
            primary_status = "partial_success"
            if parsed["scope"]["status"] not in ("present", "fallback") and parsed["national_foreword"]["status"] != "present":
                review_required = True
                warnings.append("Neither Scope nor National Foreword found")
        else:
            primary_status = "failed"
        
        # A9: Source URL provenance
        # verified_source_url is ONLY set for MATCH (not match_version_uncertain, not unknown)
        verified_source_url = fetch_result.get("preview_url") if is_verified_evidence else None
        candidate_urls = fetch_result.get("attempted_urls", [])
        if not candidate_urls and fetch_result.get("preview_url"):
            candidate_urls = [fetch_result["preview_url"]]
        
        # Gate P0-4: Raw Excel row preservation
        raw_excel = std.get("raw_excel")
        if not raw_excel:
            raw_excel = {
                "sheet_name": std.get("sheet_name", "Sheet1"),
                "row_number": std.get("physical_row_number") or (std.get("row_index", 0) + 3),
                "row_sha256": None,
                "columns": {k: str(v) for k, v in std.items() if v is not None},
            }
            
        # Gate P0-6: Non-overwriting timestamps & provenance
        first_seen = None
        if prior_rec:
            first_seen = prior_rec.get("provenance", {}).get("first_seen_at") or prior_rec.get("collected_at")
        if not first_seen:
            first_seen = datetime.now(timezone.utc).isoformat()
            
        provenance = {
            "first_seen_at": first_seen,
            "first_run_started_at": self.state.get("first_run_started_at"),
            "last_run_started_at": self.state.get("last_run_started_at"),
            "record_last_processed_at": datetime.now(timezone.utc).isoformat(),
            "parser_version": PARSER_VERSION,
            "collector_version": __version__,
            "schema_version": SCHEMA_VERSION,
            "source_file": self.input_excel.name,
            "source_sheet": std.get("sheet_name", "Sheet1"),
            "source_row": std.get("physical_row_number") or std.get("row_index", 0),
            "row_sha256": raw_excel.get("row_sha256"),
        }
            
        # 1. Standard Record (V1.2 Schema)
        record = {
            "record_version": "1.2",
            "input": {
                "source_file": self.input_excel.name,
                "source_sheet": std.get("sheet_name", "Sheet1"),
                "source_row": std.get("physical_row_number") or std.get("row_index", 0),
                "raw_standard_number": raw_num,
                "raw_title": std.get("title"),
            },
            "provenance": provenance,
            "raw_excel": raw_excel,  # Gate P0-4
            "identity": identity_dict,
            "content": {
                "header_text": parsed.get("header_text"),  # A6
                "title": title_val,
                "title_source": parsed.get("title_source", "excel"),
                "ics_raw": parsed.get("ics_raw"),
                "ics_codes": parsed.get("ics_codes", []),
                "committee": parsed.get("committee"),
                "scope": parsed["scope"],
                "national_foreword": parsed["national_foreword"],
                "notes": parsed["notes"],
                "evidence_verified": is_verified_evidence,  # A8: whether content is from a verified page
                "references_status": parsed.get("references_status", "absent"),
            },
            "references": all_refs,
            "source": {
                "verified_source_url": verified_source_url,  # A9: null unless MATCH
                "candidate_urls": candidate_urls,  # A9: all attempted
                "preview_url": fetch_result.get("preview_url"),  # backwards compat
                "requested_preview_id": fetch_result.get("requested_preview_id"),
                "matched_preview_id": fetch_result.get("matched_preview_id"),
                "attempted_urls": fetch_result.get("attempted_urls", []),
                "http_status": fetch_result.get("http_status"),
                "fetch_status": fetch_result.get("fetch_status"),
                "identity_match": identity_match,
                "page_identity_evidence": fetch_result.get("page_identity_evidence"),
            },
            "primary_status": primary_status,  # A10: non-overlapping
            "quality": {
                "manual_review_required": review_required,
                "warnings": warnings,
                "errors": errors,
            },
            # Backwards-compatible top-level keys for existing excel writers / consumers
            "is_number": identity_dict.get("standard_designation", "").split(":")[0],
            "year": identity_dict.get("year"),
            "title": title_val,
            "department": self.department,
            "ics_raw": parsed.get("ics_raw"),
            "ics_code": parsed.get("ics_code"),
            "ics_codes": parsed.get("ics_codes", []),
            "committee": parsed.get("committee"),
            "scope": parsed["scope"]["text"],
            "national_foreword": parsed["national_foreword"]["text"],
            "notes": [item["text"] for item in parsed["notes"]["items"]],
            "extraction_status": primary_status,
            "review_required": review_required,
            "source_url": verified_source_url,  # A9: only verified
            "references_count": len(all_refs),
            "notes_flag": "; ".join(warnings) if warnings else None,
            "collected_at": datetime.now(timezone.utc).isoformat(),
        }
        
        # 2. Source Availability Record
        avail_record = {
            "standard_designation": canonical_id,
            "original_standard_number": raw_num,
            "excel_present": True,
            "preview": {
                "attempted": bool(candidates),
                "found": fetch_result.get("content_valid", False),
                "identity": identity_match,
                "preview_id": fetch_result.get("matched_preview_id"),
                "verified_source_url": verified_source_url,
                "preview_url": fetch_result.get("preview_url"),
            },
            "scope": parsed["scope"]["status"],
            "national_foreword": parsed["national_foreword"]["status"],
            "notes": parsed["notes"]["status"],
            "references": "present" if len(all_refs) > 0 else "absent",
            "reference_counts": {
                "total": len(all_refs),
                "formal_rows": sum(1 for r in all_refs if r.get("relationship") != "dual_numbering"),
                "extracted_references": len(all_refs),
                "is": sum(1 for r in all_refs if r["target"]["family"] in ("IS", "IS/IEC", "IS/ISO", "IS/ISO/IEC")),
                "international": sum(1 for r in all_refs if r["target"]["family"] in ("ISO", "IEC", "ISO/IEC", "ISO/IEC GUIDE")),
                "sp": sum(1 for r in all_refs if r["target"]["family"] == "SP"),
                "dual_numbered": sum(1 for r in all_refs if r.get("relationship") == "dual_numbering"),
                "continuation_resolved": sum(1 for r in all_refs if r.get("evidence", {}).get("continuation_resolved")),
            },
            "primary_status": primary_status,
            "evidence_verified": is_verified_evidence,
            "manual_review_required": review_required,
            "warnings": warnings,
        }
        
        return record, all_refs, avail_record
        
    def _fetch_from_cache_only(self, candidate_ids: list[str], requested_identity: dict = None) -> dict:
        """Fetch strictly from local raw cache (zero network calls). Supports V1.2 identity statuses."""
        attempted_urls = []
        version_uncertain_result = None
        
        for cand_id in candidate_ids:
            url = get_preview_url(cand_id)
            attempted_urls.append(url)
            html = self.fetcher.get_cached_html(cand_id)
            if not html:
                continue
                
            id_status, reason, evidence = self.fetcher.verify_page_identity(requested_identity, html)
            if id_status in ("match_version_verified", "match"):
                return {
                    "requested_preview_id": candidate_ids[0],
                    "matched_preview_id": cand_id,
                    "preview_url": url,
                    "attempted_urls": attempted_urls,
                    "http_status": 200,
                    "fetch_status": "success",
                    "identity_match": "match_version_verified",
                    "page_identity_evidence": evidence,
                    "html": html,
                    "cache_path": str(self.fetcher._get_cache_path(cand_id)),
                    "content_valid": True,
                    "error": None,
                    "warnings": [],
                }
            elif id_status in ("match_version_uncertain", "match_identity_only_version_uncertain"):
                if not version_uncertain_result:
                    version_uncertain_result = {
                        "requested_preview_id": candidate_ids[0],
                        "matched_preview_id": cand_id,
                        "preview_url": url,
                        "attempted_urls": attempted_urls,
                        "http_status": 200,
                        "fetch_status": "review_required",
                        "identity_match": id_status,
                        "page_identity_evidence": evidence,
                        "html": html,
                        "cache_path": str(self.fetcher._get_cache_path(cand_id)),
                        "content_valid": True,
                        "error": None,
                        "warnings": [f"Cached page version uncertain ({id_status}): {reason}"],
                    }
                continue
            elif id_status == "mismatch":
                continue
            else:  # unknown
                return {
                    "requested_preview_id": candidate_ids[0],
                    "matched_preview_id": cand_id,
                    "preview_url": url,
                    "attempted_urls": attempted_urls,
                    "http_status": 200,
                    "fetch_status": "review_required",
                    "identity_match": "unknown",
                    "page_identity_evidence": evidence,
                    "html": html,
                    "cache_path": str(self.fetcher._get_cache_path(cand_id)),
                    "content_valid": True,
                    "error": None,
                    "warnings": [f"Cached page identity unknown: {reason}"],
                }
        
        if version_uncertain_result:
            return version_uncertain_result
                
        return {
            "requested_preview_id": candidate_ids[0] if candidate_ids else None,
            "matched_preview_id": None,
            "preview_url": None,
            "attempted_urls": attempted_urls,
            "http_status": None,
            "fetch_status": "failed",
            "identity_match": "unknown",
            "page_identity_evidence": None,
            "html": None,
            "cache_path": None,
            "content_valid": False,
            "error": "No matching cached HTML found for candidates",
            "warnings": ["Cached preview absent"],
        }

    def _consolidate_output_files(self, original_headers: list[str], original_rows: list[list], active_ids: set[str] = None):
        """
        Consolidate in-memory records to final JSONL, CSV, and Excel files atomically.
        Guarantees zero duplicate canonical records.
        Gate P0-3: Snapshot Mode prunes records that are no longer in the active input workbook.
        """
        if active_ids is None or self.incremental:
            target_records = list(self.records_by_id.values())
            target_refs = list(self.references_by_key.values())
            target_avails = list(self.availability_by_id.values())
        else:
            target_records = [r for desig, r in self.records_by_id.items() if desig in active_ids]
            target_refs = [r for (src, tgt, rel), r in self.references_by_key.items() if src in active_ids]
            target_avails = [a for desig, a in self.availability_by_id.items() if desig in active_ids]
            
            pruned_count = len(self.records_by_id) - len(target_records)
            if pruned_count > 0:
                logger.info(f"Snapshot Mode: pruned {pruned_count} historical records not present in current input workbook")

        # 1. Write standards.jsonl atomically
        tmp_std = self.standards_jsonl.with_suffix(".tmp")
        with open(tmp_std, "w", encoding="utf-8") as f:
            for rec in target_records:
                f.write(json.dumps(rec, ensure_ascii=False, default=str) + "\n")
        tmp_std.replace(self.standards_jsonl)
        
        # 2. Write references.jsonl atomically
        tmp_ref = self.references_jsonl.with_suffix(".tmp")
        with open(tmp_ref, "w", encoding="utf-8") as f:
            for ref in target_refs:
                f.write(json.dumps(ref, ensure_ascii=False, default=str) + "\n")
        tmp_ref.replace(self.references_jsonl)
        
        # 3. Write source_availability.jsonl atomically
        tmp_avail = self.availability_jsonl.with_suffix(".tmp")
        with open(tmp_avail, "w", encoding="utf-8") as f:
            for avail in target_avails:
                f.write(json.dumps(avail, ensure_ascii=False, default=str) + "\n")
        tmp_avail.replace(self.availability_jsonl)
        
        # 4. Write source_availability.csv atomically
        tmp_csv = self.availability_csv.with_suffix(".tmp")
        with open(tmp_csv, "w", newline="", encoding="utf-8") as f:
            fieldnames = [
                "standard_designation", "original_standard_number", "excel_present",
                "preview_found", "preview_identity", "preview_id", "verified_source_url",
                "scope", "national_foreword", "notes", "references",
                "total_refs", "is_refs", "intl_refs", "sp_refs", "dual_numbered_refs",
                "primary_status", "evidence_verified", "manual_review_required"
            ]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for a in target_avails:
                writer.writerow({
                    "standard_designation": a.get("standard_designation"),
                    "original_standard_number": a.get("original_standard_number"),
                    "excel_present": a.get("excel_present"),
                    "preview_found": a.get("preview", {}).get("found"),
                    "preview_identity": a.get("preview", {}).get("identity"),
                    "preview_id": a.get("preview", {}).get("preview_id"),
                    "verified_source_url": a.get("preview", {}).get("verified_source_url"),
                    "scope": a.get("scope"),
                    "national_foreword": a.get("national_foreword"),
                    "notes": a.get("notes"),
                    "references": a.get("references"),
                    "total_refs": a.get("reference_counts", {}).get("total", 0),
                    "is_refs": a.get("reference_counts", {}).get("is", 0),
                    "intl_refs": a.get("reference_counts", {}).get("international", 0),
                    "sp_refs": a.get("reference_counts", {}).get("sp", 0),
                    "dual_numbered_refs": a.get("reference_counts", {}).get("dual_numbered", 0),
                    "primary_status": a.get("primary_status"),
                    "evidence_verified": a.get("evidence_verified"),
                    "manual_review_required": a.get("manual_review_required"),
                })
        tmp_csv.replace(self.availability_csv)
        
        # 5. Write enriched Excel
        self._write_enriched_excel(target_records, original_headers, original_rows)
        
    def _write_enriched_excel(self, records: list[dict], original_headers: list[str], original_rows: list[list]):
        """Write enriched Excel workbook."""
        wb = Workbook()
        ws = wb.active
        ws.title = f"{self.department} Standards"
        
        new_cols = [
            "Canonical Designation", "Scope", "Scope Status", "National Foreword", "Notes",
            "ICS Code(s)", "Committee", "References Count", "Dual-Numbered Refs",
            "Verified Source URL", "Identity Match", "Evidence Verified",
            "Primary Status", "Review Required", "Quality Warnings", "Collected At"
        ]
        
        all_headers = original_headers + new_cols
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_idx, header in enumerate(all_headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            
        lookup = {}
        for rec in records:
            desig = rec.get("identity", {}).get("standard_designation") or rec.get("is_number")
            raw_orig = rec.get("input", {}).get("raw_standard_number")
            if desig:
                lookup[desig] = rec
            if raw_orig:
                lookup[raw_orig] = rec
                
        # Find column index for standard number
        std_idx = 0
        for idx, h in enumerate(original_headers):
            if h and any(k in h.lower() for k in ["standard number", "is number", "std no"]):
                std_idx = idx
                break
                
        for row_num, row_data in enumerate(original_rows):
            excel_row = row_num + 2
            for col_idx, val in enumerate(row_data):
                ws.cell(row=excel_row, column=col_idx + 1, value=val)
                
            raw_std = str(row_data[std_idx]).strip() if std_idx < len(row_data) and row_data[std_idx] else ""
            rec = lookup.get(raw_std)
            if not rec:
                ident = parse_standard_identity(raw_std)
                if ident:
                    rec = lookup.get(ident["standard_designation"])
                    
            base_col = len(original_headers) + 1
            if rec:
                notes_items = [item["text"] for item in rec.get("content", {}).get("notes", {}).get("items", [])]
                dual_refs = sum(1 for r in rec.get("references", []) if r.get("relationship") == "dual_numbering")
                ics_val = ", ".join(rec.get("ics_codes", [])) if rec.get("ics_codes") else rec.get("ics_code")
                
                ws.cell(row=excel_row, column=base_col, value=rec.get("identity", {}).get("standard_designation"))
                ws.cell(row=excel_row, column=base_col + 1, value=rec.get("content", {}).get("scope", {}).get("text"))
                ws.cell(row=excel_row, column=base_col + 2, value=rec.get("content", {}).get("scope", {}).get("status"))
                ws.cell(row=excel_row, column=base_col + 3, value=rec.get("content", {}).get("national_foreword", {}).get("text"))
                ws.cell(row=excel_row, column=base_col + 4, value=" | ".join(notes_items) if notes_items else None)
                ws.cell(row=excel_row, column=base_col + 5, value=ics_val)
                ws.cell(row=excel_row, column=base_col + 6, value=rec.get("committee"))
                ws.cell(row=excel_row, column=base_col + 7, value=len(rec.get("references", [])))
                ws.cell(row=excel_row, column=base_col + 8, value=dual_refs)
                ws.cell(row=excel_row, column=base_col + 9, value=rec.get("source", {}).get("verified_source_url"))
                ws.cell(row=excel_row, column=base_col + 10, value=rec.get("source", {}).get("identity_match"))
                ws.cell(row=excel_row, column=base_col + 11, value=str(rec.get("content", {}).get("evidence_verified", False)))
                ws.cell(row=excel_row, column=base_col + 12, value=rec.get("primary_status"))
                ws.cell(row=excel_row, column=base_col + 13, value=str(rec.get("quality", {}).get("manual_review_required", False)))
                ws.cell(row=excel_row, column=base_col + 14, value="; ".join(rec.get("quality", {}).get("warnings", [])))
                ws.cell(row=excel_row, column=base_col + 15, value=rec.get("collected_at"))
                
        tmp_excel = self.enriched_excel.with_suffix(".tmp")
        try:
            wb.save(tmp_excel)
            tmp_excel.replace(self.enriched_excel)
            logger.info(f"Enriched Excel written to {self.enriched_excel}")
        except PermissionError:
            logger.warning(f"Could not overwrite {self.enriched_excel} (file open in another program). Saved as {tmp_excel}")

    def _generate_report(self, standards: list[dict], run_start: datetime, run_end: datetime) -> dict:
        """Generate run-level collection report per V1.2 Data Contract."""
        records = list(self.records_by_id.values())
        refs = list(self.references_by_key.values())
        
        # A10: Use primary_status for counts (non-overlapping)
        success_count = sum(1 for r in records if r.get("primary_status") == "success")
        partial_count = sum(1 for r in records if r.get("primary_status") == "partial_success")
        failed_count = sum(1 for r in records if r.get("primary_status") == "failed")
        review_count = sum(1 for r in records if r.get("quality", {}).get("manual_review_required"))
        
        scope_present = sum(1 for r in records if r.get("content", {}).get("scope", {}).get("status") == "present")
        scope_fallback = sum(1 for r in records if r.get("content", {}).get("scope", {}).get("status") == "fallback")
        foreword_found = sum(1 for r in records if r.get("content", {}).get("national_foreword", {}).get("status") == "present")
        notes_found = sum(1 for r in records if r.get("content", {}).get("notes", {}).get("status") == "present")
        
        evidence_verified = sum(1 for r in records if r.get("content", {}).get("evidence_verified"))
        version_uncertain = sum(
            1 for r in records
            if r.get("source", {}).get("identity_match") in ("match_version_uncertain", "match_identity_only_version_uncertain")
        )
        
        stds_with_refs = len(set(r.get("source_standard") for r in refs if r.get("source_standard")))
        intl_refs = sum(1 for r in refs if r.get("target", {}).get("family") in ("ISO", "IEC", "ISO/IEC", "ISO/IEC GUIDE"))
        sp_refs = sum(1 for r in refs if r.get("target", {}).get("family") == "SP")
        dual_refs = sum(1 for r in refs if r.get("relationship") == "dual_numbering")
        
        first_started = self.state.get("first_run_started_at") or self.state.get("run_started_at") or run_start.isoformat()
        return {
            "run_id": run_start.isoformat(),
            "first_run_started_at": first_started,
            "last_run_started_at": run_start.isoformat(),
            "run_completed_at": run_end.isoformat(),
            "department": self.department,
            "input_file": self.input_excel.name,
            "input_rows": len(standards),
            "run_manifest": self.state.get("run_manifest", {}),
            "statistics": {
                "input_rows_count": len(standards),
                "attempted_count": len(standards) - self.stats.get("skipped_cached", 0),
                "skipped_cache_count": self.stats.get("skipped_cached", 0),
                "reparsed_count": self.stats.get("reparsed", 0),
                "output_snapshot_count": len(records),
            },
            "standards": {
                "processed": len(records),
                "success": success_count,
                "partial_success": partial_count,
                "failed": failed_count,
                "review_required": review_count,
                "note": "review_required is a quality flag, not a status. success+partial+failed=processed."
            },
            "evidence": {
                "verified": evidence_verified,
                "version_uncertain": version_uncertain,
                "note": "version_uncertain pages have content but years differ from requested."
            },
            "content": {
                "scope_present": scope_present,
                "scope_fallback": scope_fallback,
                "scope_found": scope_present + scope_fallback,
                "foreword_found": foreword_found,
                "notes_found": notes_found,
            },
            "references": {
                "standards_with_references": stds_with_refs,
                "total_reference_records": len(refs),
                "formal_reference_rows": len(refs) - dual_refs,
                "extracted_references": len(refs),
                "international_references": intl_refs,
                "sp_references": sp_refs,
                "dual_numbered_references": dual_refs,
                "continuation_resolved": sum(1 for r in refs if r.get("evidence", {}).get("continuation_resolved")),
            },
            "output_paths": {
                "standards_jsonl": str(self.standards_jsonl),
                "references_jsonl": str(self.references_jsonl),
                "source_availability_jsonl": str(self.availability_jsonl),
                "source_availability_csv": str(self.availability_csv),
                "enriched_excel": str(self.enriched_excel),
            }
        }
